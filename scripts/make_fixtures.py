"""Generates the fictional FAP memos used by the tests (python-docx required).

All agencies, people, figures and identifiers in these documents are invented.
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "Tests/ReviewCoreTests/Fixtures")
OUT.mkdir(parents=True, exist_ok=True)


def header_block(doc, title):
    t = doc.add_paragraph()
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, value in [
        ("TO:", "Office of Fleet and Asset Management (OFAM)"),
        ("FROM:", "Dana Whitfield, Fleet Coordinator, Department of Coastal Surveys"),
        ("DATE:", "March 4, 2026"),
        ("SUBJECT:", "Fleet Acquisition Plan – Request for Two Replacement Field Vehicles"),
    ]:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
        p.add_run("\t" + value)


def basic_memo():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    header_block(doc, "MEMORANDUM")
    doc.add_heading("Background", level=1)
    p = doc.add_paragraph("The Department of Coastal Surveys (DCS) requests approval to replace two ")
    r = p.add_run("2011 Ford Ranger")
    r.italic = True
    p.add_run(" pickup trucks that have exceeded their useful life. Both vehicles is used daily by survey crews "
              "traveling to remote shoreline monitoring stations.")
    doc.add_heading("Justification", level=1)
    p = doc.add_paragraph("The department’s current fleet averaged ")
    r = p.add_run("8,200 miles")
    r.bold = True
    p.add_run(" per vehicle last fiscal year, which we believe demonstrates “high utilization” of the existing fleet.")
    doc.add_paragraph("The replacement vehicles will be gasoline-powered because charging infrastructure is not "
                      "available at all monitoring stations. An exemption from the ZEV mandate is therefore requested.")
    p = doc.add_paragraph("Additional detail is available on the ")
    add_hyperlink(p, "https://example.org/dcs-fleet", "DCS fleet page")
    p.add_run(". Funding has been identified in the current budget.")
    doc.add_heading("Vehicle Summary", level=1)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i, row in enumerate([("Unit", "Model year", "Annual miles"),
                             ("DCS-0142", "2011", "7,950"),
                             ("DCS-0178", "2011", "8,450")]):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    doc.add_paragraph("Please contact the fleet coordinator with any questions regarding this request.")
    doc.save(OUT / "basic-memo.docx")


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def fragment_runs(paragraph, cuts):
    """Splits a single-run paragraph at the given offsets the way Word's rsid tracking does."""
    run = paragraph.runs[0]
    text = run.text
    pieces, last = [], 0
    for c in cuts + [len(text)]:
        pieces.append(text[last:c])
        last = c
    run.text = pieces[0]
    anchor = run._r
    for i, piece in enumerate(pieces[1:], start=1):
        new = copy.deepcopy(run._r)
        new.set(qn("w:rsidR"), f"00A{i:05d}")
        for t in new.findall(qn("w:t")):
            t.text = piece
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        anchor.addnext(new)
        anchor = new


def complex_memo():
    doc = Document()
    header_block(doc, "FLEET ACQUISITION PLAN MEMO")

    p = doc.add_paragraph("Utilization for the requested vehicles averaged 6,100 miles annually, well below the threshold.")
    fragment_runs(p, [4, 15, 38, 41, 60])

    # A complex field (DATE) in the middle of body text: its result must stay untouched.
    p = doc.add_paragraph("This plan was printed on ")
    add_field(p, " DATE \\@ \"MMMM d, yyyy\" ", "October 10, 2026")
    p.add_run(" and supersedes the prior submission.")

    # A footnote reference run.
    p = doc.add_paragraph("Mileage figures come from the department’s telematics report")
    add_footnote_reference(doc, p, "Telematics export dated February 2026 (fictional).")
    p.add_run(" covering the last twelve months.")

    # An existing tracked change from another reviewer.
    p = doc.add_paragraph("The department will dispose of the replaced units through ")
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "901")
    ins.set(qn("w:author"), "Earlier Reviewer")
    ins.set(qn("w:date"), "2026-03-01T09:00:00Z")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "state surplus "
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    ins.append(r)
    p._p.append(ins)
    p.add_run("auction after delivery of the new vehicles.")

    # A content control (structured document tag) holding a form value.
    p = doc.add_paragraph("Requested funding: ")
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "Amount")
    sdt_pr.append(alias)
    sdt.append(sdt_pr)
    content = OxmlElement("w:sdtContent")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "$96,400"
    r.append(t)
    content.append(r)
    sdt.append(content)
    p._p.append(sdt)
    p.add_run(" from the Motor Vehicle Account.")

    # An existing comment from someone else.
    p = doc.add_paragraph("Both vehicles will be assigned to the Eureka field office.")
    doc.add_comment(p.runs, text="Confirm the office location.", author="Earlier Reviewer", initials="ER")

    p = doc.add_paragraph("Repeated phrase: the vehicle, the vehicle, the vehicle.")
    p = doc.add_paragraph("Line one of an address block")
    p.add_run().add_break()
    p.add_run("Line two after a manual line break.")
    doc.save(OUT / "complex-memo.docx")


def add_field(paragraph, instruction, result):
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r
    p = paragraph._p
    p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = instruction
    r.append(it)
    p.append(r)
    p.append(fld("separate"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = result
    r.append(t)
    p.append(r)
    p.append(fld("end"))


def add_footnote_reference(doc, paragraph, text):
    """Adds a footnotes part with one footnote and references it from the paragraph."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    xml = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {ns}>'
           '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
           '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
           f'<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
           f'<w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:footnote></w:footnotes>')
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                xml.encode(), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "superscript")
    rpr.append(va)
    r.append(rpr)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "1")
    r.append(ref)
    paragraph._p.append(r)


basic_memo()
complex_memo()
print("wrote fixtures to", OUT)
